import openpyxl, json, pandas as pd
P=json.load(open('prep.json'))
wb=openpyxl.Workbook(); ws=wb.active; ws.title='CAIF recursos'
yrs=P['caif_years']; ws.append(['Concepto']+yrs)
keep=['1) Ingresos corrientes','Ingresos tributarios','Ingresos no tributarios','Venta de bienes y servicios','Rentas de la propiedad','Transferencias corrientes','4) Recursos de capital','Recursos propios de capital','Transferencias de capital','Disminución de la inversión financiera','6) Recursos totales (1+4)']
seen=set()
for c in P['caif']:
    n=c['concepto'].strip()
    if n in keep and n not in seen: seen.add(n); ws.append([n]+[c['vals'].get(y) for y in yrs])
ws.append([]); ws.append(['Millones de pesos corrientes. Fuente: IDECBA SP_Fi_AX01 (2025 provisorio); 2027P: Planilla 16 Proyecto 2027.'])
w2=wb.create_sheet('constante ($2027)'); w2.append(['Concepto']+yrs)
for r in list(ws.iter_rows(min_row=2,max_row=1+len(seen),values_only=True)):
    w2.append([r[0]]+[(v*P['deflator']['p2027'][y] if v is not None else None) for v,y in zip(r[1:],yrs)])
w2.append([]); w2.append([P['deflator']['note']])
rc=pd.read_pickle('reca.pkl'); w3=wb.create_sheet('recaudacion mensual'); w3.append(list(rc.columns))
for r in rc.itertuples(index=False): w3.append(list(r))
w3.append(['Millones de pesos corrientes. Fuente: IDECBA SP_IP_Reca_concepto (AGIP), ene-1997 a ago-2026.'])
s=openpyxl.load_workbook('raw/SP_IP_ISIB_por_Regimen.xlsx').worksheets[0]; w4=wb.create_sheet('IIBB por regimen')
for r in s.iter_rows(values_only=True): w4.append(list(r))
D=json.load(open('data.json')); w5=wb.create_sheet('recursos 2027 rubros'); w5.append(['Rubro','Millones de pesos 2027'])
for r in D['rec2027']: w5.append([r[0],r[1]*1000])
w6=wb.create_sheet('deflactores'); w6.append(['año','factor a $ hoy (sep-26)','factor a $ 2027','inflación promedio anual'])
for y in sorted(P['deflator']['hoy']): w6.append([int(y),P['deflator']['hoy'][y],P['deflator']['p2027'][y],P['deflator']['infl_avg'].get(y)])
wb.save('out/recursos_CABA_CAIF_y_aperturas.xlsx'); print('ok')
