import pandas as pd
b=pd.read_pickle('base_gasto.pkl')
b.to_csv('/home/claude/out/base_gasto_CABA_trimestral.csv.gz',index=False,compression='gzip')
from openpyxl import Workbook
from openpyxl.styles import Font
wb=Workbook(write_only=True)
ws=wb.create_sheet('base')
ws.append(list(b.columns))
for r in b.itertuples(index=False): ws.append(list(r))
n=wb.create_sheet('notas')
for line in [
 ['Base de gasto CABA - ejecución trimestral acumulada (pesos corrientes)'],
 ['Clave: año, trim, carácter, jurisdicción, programa, finalidad, función, inciso, principal, parcial, subparcial, fuente de financiamiento'],
 ['Valores acumulados al cierre de cada trimestre (sanción, vigente, definitivo, devengado), en pesos.'],
 ['Fuente trimestral: data.buenosaires.gob.ar (Presupuesto ejecutado) - CSV del CDN de la Ciudad, descargados 12/10/2026.'],
 ['2020-T4, 2021-T4, 2023-T4 y 2024-T4: cierre anual de la base propia (Sheet1 de consolidado). El CSV 2024-T4 publicado está incompleto (faltan partidas de Seguridad, Movilidad y Desarrollo Humano) y el 2023-T4 tiene vigente duplicado.'],
 ['Se eliminaron filas de subtotales sin códigos presentes en 2022-T3, 2024-T1 y 2024-T2 (duplicaban el total).'],
 ['No publicados en el portal: 2020-T2/T3, 2021-T1/T3.'],
 ['2023-T1: el crédito sancionado de la jurisdicción 26 aparece incompleto en el archivo fuente.'],
 ['Programas: códigos por jurisdicción; las descripciones corresponden a cada año.'],
]: n.append(line)
wb.save('/home/claude/out/base_gasto_CABA_trimestral.xlsx')
print('ok')
