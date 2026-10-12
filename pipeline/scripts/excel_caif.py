import openpyxl, json
from openpyxl.utils import get_column_letter as L
P=json.load(open('prep.json'))
wb=openpyxl.load_workbook('/home/claude/in/58287fbc-CAIFdeuda.xlsx')
ws=wb['dataset- CAIF corriente']
for mr in list(ws.merged_cells.ranges): ws.unmerge_cells(str(mr))
src=openpyxl.load_workbook('raw/SP_Fi_AX01.xlsx').worksheets[0]
# clear and copy values
for row in ws.iter_rows(min_row=1,max_row=60,max_col=40):
    for c in row: c.value=None
r_out=1
for r in src.iter_rows(values_only=True):
    for j,v in enumerate(r): ws.cell(r_out+1,j+1).value=v
    r_out+=1
# add 2027P column at AG (col 33)
ws.cell(3,33).value='2027P'
caif={c['concepto']:c['vals'] for c in P['caif']}
for i in range(4,41):
    name=ws.cell(i,1).value
    if not name: continue
    # match by row position via P['caif'] order
for idx,c in enumerate(P['caif']):
    ws.cell(4+idx,33).value=c['vals'].get('2027P')
ws.cell(2,33).value='Proyecto 2027 (Planilla 16)'
# ipc sheet: extend
ip=wb['ipc']
rates={2025:None,2026:None,2027:0.18}
infl=P['deflator']['infl_avg']
# use IPCBA dec/dec for 2025-2026 (2026 with oct-dic at 2%)
import openpyxl as ox
m={}
for r in ox.load_workbook('raw/ipcba.xlsx',data_only=True).worksheets[0].iter_rows(values_only=True):
    if hasattr(r[0],'year') and isinstance(r[1],(int,float)): m[(r[0].year,r[0].month)]=r[1]
for k in (10,11,12): m[(2026,k)]=m[(2026,9)]*1.02**(k-9)
rates[2025]=m[(2025,12)]/m[(2024,12)]-1; rates[2026]=m[(2026,12)]/m[(2025,12)]-1
rowmap={ip.cell(r,1).value:r for r in range(2,ip.max_row+1)}
last=max(rowmap.values())
for y in (2025,2026,2027):
    r=rowmap.get(y) or (last:=last+1)
    ip.cell(r,1).value=y; ip.cell(r,2).value=rates[y]; ip.cell(r,3).value=f'=B{r}+1'
for r in range(2,last+1): ip.cell(r,4).value=f'=PRODUCT(C{r+1}:$C${last})' if r<last else 1
ip.cell(1,8).value='Nota: 2025-2026 = IPCBA dic/dic (oct-dic 2026 estimado al 2% mensual); 2027 = pauta 18% del proyecto. Base de coefacum: diciembre 2027.'
# constante sheet: extend formulas to 2025 (AF) and 2027P (AG)
cs=wb['constante']
cs.cell(1,32).value=2025; cs.cell(1,33).value=2027
for r in range(2,39):
    f=cs.cell(r,31).value
    if isinstance(f,str) and f.startswith('='):
        for col,srccol in ((32,'AF'),(33,'AG')):
            cs.cell(r,col).value=f.replace("'dataset- CAIF corriente'!AE","'dataset- CAIF corriente'!"+srccol).replace('AE$1',L(col)+'$1')
# new raw sheets
for fn,name in [('raw/SP_GP_AX01.xlsx','IDECBA gasto finalidad'),('raw/SP_SD_AX01.xlsx','IDECBA stock deuda 2025'),('raw/SP_TP_AX01.xlsx','IDECBA titulos 2025')]:
    s=openpyxl.load_workbook(fn).worksheets[0]; n=wb.create_sheet(name)
    for r in s.iter_rows(values_only=True): n.append(list(r))
wb.save('out/CAIF+deuda_actualizado.xlsx'); print('ok', rates)
