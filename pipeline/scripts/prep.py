import openpyxl, pandas as pd, numpy as np, json, re, sys
sys.path.insert(0,'scripts')
from planillas import parse_rot, d as PD
OUT={}
# ---------- IPCBA deflators
wb=openpyxl.load_workbook('raw/ipcba.xlsx',data_only=True); ws=wb.worksheets[0]
ipc={}
for r in ws.iter_rows(values_only=True):
    if hasattr(r[0],'year') and isinstance(r[1],(int,float)): ipc[(r[0].year,r[0].month)]=r[1]
last=max(ipc); 
# project Oct-Dec 2026 at 2.0% monthly
y,m=last; v=ipc[last]
while (y,m)<(2026,12):
    m+=1
    if m>12: y+=1;m=1
    v*=1.02; ipc[(y,m)]=v
ann={yy:np.mean([ipc[(yy,k)] for k in range(1,13)]) for yy in range(2013,2027)}
wbu=openpyxl.load_workbook('/home/claude/in/58287fbc-CAIFdeuda.xlsx',data_only=True)
nat={r[0]:r[1] for r in wbu['ipc'].iter_rows(min_row=2,values_only=True) if isinstance(r[0],int)}
for yy in range(2013,1994,-1): ann[yy-1]=ann[yy]/(1+nat[yy])
HOY=ipc[(2026,9)]; PAUTA=0.18; REM=0.2154
ytd={}
for yy in range(2013,2027):
    for q in range(1,5): ytd[(yy,q)]=np.mean([ipc[(yy,k)] for k in range(1,3*q+1)])
OUT['deflator']={
 'hoy':{str(k):HOY/v for k,v in sorted(ann.items())},
 'hoy_ytd':{f'{k[0]}-{k[1]}':HOY/v for k,v in sorted(ytd.items())},
 'p2027':{**{str(k):HOY*1.02**3/v*(1+PAUTA) for k,v in sorted(ann.items())},'2027P':1.0},
 'p2027_ytd':{f'{k[0]}-{k[1]}':HOY*1.02**3/v*(1+PAUTA) for k,v in sorted(ytd.items())},
 'infl_avg':{str(k):ann[k]/ann[k-1]-1 for k in range(1996,2027)},
 'note':'Pesos de hoy = IPCBA sep-2026. Promedios anuales IPCBA (IDECBA) desde 2013; 1995-2012 empalme con la serie alternativa de la base CEER (IPC provincias). Oct-dic 2026 proyectado al 2% mensual. Pesos 2027 = pesos de hoy (sep-2026) x 1,02^3 (oct-dic 2026 al 2% mensual) x 1,18 (pauta de inflacion 2027 del proyecto).'}
# ---------- CAIF
wb=openpyxl.load_workbook('raw/SP_Fi_AX01.xlsx',data_only=True); ws=wb.worksheets[0]
rows=list(ws.iter_rows(values_only=True))
hdr=rows[1]; years=[int(str(h).replace('*','')) for h in hdr[1:] if h]
caif=[]
for r in rows[2:]:
    if r[0] and any(isinstance(x,(int,float)) for x in r[1:]):
        caif.append({'concepto':r[0].strip(),'vals':{str(y):(r[i+1] if isinstance(r[i+1],(int,float)) else None) for i,y in enumerate(years)}})
# 2027 proyecto from Planilla 16 (millones)
p16={'1) Ingresos corrientes':24092481.443901,'Ingresos tributarios':19094214.003433,'Ingresos no tributarios':663825.629989,'Venta de bienes y servicios':365617.417614,'Rentas de la propiedad':951290.901460,'Transferencias corrientes':3017533.491405,
'2) Gastos corrientes':19176736.173780,'Remuneraciones al personal':9950480.728332,'Gastos de consumo':6428377.453576,'Transferencias corrientes (gasto)':2797877.991872,
'3) Resultado económico (1-2)':4915745.270121,'4) Recursos de capital':3175.598228,'Recursos propios de capital':0,'Transferencias de capital':0,'Disminución de la inversión financiera':3175.598228,
'5) Gastos de capital':4325717.119862,'Inversión real directa':4317027.252548,'Transferencias de capital (gasto)':4547.521394,'Inversión financiera':4142.345920,
'6) Recursos totales (1+4)':24095657.042129,'7) Gasto primario (2+5)':23502453.293642,'8) Resultado primario (6-7)':593203.748487,'9) Intereses de la deuda pública':591887.305621,'10) Gastos totales (7+9)':24094340.599263,'11) Resultado financiero antes de figurativas (6-10)':1316.442866}
# assign by order/name
seen_tc=0; seen_tk=0
for c in caif:
    n=c['concepto']; key=None
    nn=n.strip().lower()
    for k in p16:
        if k.lower().startswith(nn[:18]) : key=k;break
    if nn.startswith('transferencias corrientes'):
        key='Transferencias corrientes' if seen_tc==0 else 'Transferencias corrientes (gasto)'; seen_tc+=1
    if nn.startswith('transferencias de capital'):
        key='Transferencias de capital' if seen_tk==0 else 'Transferencias de capital (gasto)'; seen_tk+=1
    if nn.startswith('disminución de la inversión financiera') and 'Disminución' in str(c.get('_x','')): pass
    c['vals']['2027P']=p16.get(key) if key else None
OUT['caif']=caif; OUT['caif_years']=[str(y) for y in years]+['2027P']
# ---------- Planillas 2027
cols,rws,data=parse_rot(5); OUT['p2027_jur_inc']={r:v for r,v in data.items()}
cols,rws,data=parse_rot(4); OUT['p2027_fun_inc']={r:v for r,v in data.items()}
json.dump(OUT,open('prep.json','w'),ensure_ascii=False,default=float)
for c in caif: print(c['concepto'][:45].ljust(45), c['vals'].get('2024'), c['vals'].get('2025'), c['vals'].get('2027P'))
print({k:round(v,3) for k,v in OUT['deflator']['infl_avg'].items() if int(k)>2005}); print(round(OUT['deflator']['hoy']['2025'],3), round(OUT['deflator']['p2027']['2026'],3))
